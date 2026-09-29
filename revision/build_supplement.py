"""Assemble the Supplementary Material (.docx) from the methods/results drafts, analysis CSVs and figures.

Run after the analysis scripts. Output: revision/outputs/Supplementary_Material_revision.docx
"""
import re

import pandas as pd
from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from common import BASELINE, FIG_DIR, MODELS, OUT

SOURCES = MODELS + [BASELINE]
SHORT_SOURCE = {BASELINE: "Indep. baseline"}
MAX_FIG_HEIGHT = 7.8  # inches

SAFE_HARBOR = [
    ("1. Names", "Patient initials in source data", "Excluded from training; never modeled"),
    ("2. Geographic subdivisions smaller than a state", "None collected", "—"),
    ("3. All elements of dates (except year); ages > 89",
     "Date of birth, date of surgery; age at diagnosis and at surgery",
     "Dates excluded from training. Remaining time variables are durations (survival, follow-up, "
     "length of stay), not dates. Maximum age 79 in real and synthetic data (no ages > 89)"),
    ("4. Telephone numbers", "None collected", "—"),
    ("5. Fax numbers", "None collected", "—"),
    ("6. Email addresses", "None collected", "—"),
    ("7. Social Security numbers", "None collected", "—"),
    ("8. Medical record numbers", "MRN in source data", "Excluded from training; never modeled"),
    ("9. Health plan beneficiary numbers", "None collected", "—"),
    ("10. Account numbers", "None collected", "—"),
    ("11. Certificate/license numbers", "None collected", "—"),
    ("12. Vehicle identifiers", "None collected", "—"),
    ("13. Device identifiers and serial numbers", "None collected", "—"),
    ("14. Web URLs", "None collected", "—"),
    ("15. IP addresses", "None collected", "—"),
    ("16. Biometric identifiers", "None collected", "—"),
    ("17. Full-face photographs", "None collected", "—"),
    ("18. Any other unique identifying number, characteristic or code",
     "No free-text fields (≤ 6 categories per variable); rare combinations of clinical characteristics",
     "Combinations of clinical characteristics evaluated with attribute inference and membership analyses "
     "(Supplementary Methods S5–S6)"),
]

FIGURES = [
    ("S_missingness", "Supplementary Figure S1. Variable-level missingness before imputation (variables with ≥ 1 "
     "missing value; labels give the number of missing patients out of 18)."),
    ("S_fidelity_by_seed", "Supplementary Figure S2. SDMetrics quality scores across 10 independently seeded runs "
     "per synthesizer. Points are individual runs; bars are means. The independent-marginals baseline keeps "
     "each variable's distribution but has no relationships between variables."),
    ("S_utility_associations", "Supplementary Figure S3. Prespecified clinical associations in the real cohort "
     "(black, log odds ratio with 95% CI) and in synthetic data (mean and range across 10 runs)."),
    ("S_utility_global_structure", "Supplementary Figure S4. Preservation of the global correlation structure. "
     "Left: correlation between the 1,485 real and synthetic pairwise Spearman correlations. Right: sign "
     "agreement among strong real correlations (|ρ| ≥ 0.3). Points are individual runs; bars are means."),
    ("S_privacy_original_method_evaluable", "Supplementary Figure S5. Original attack specification: number of "
     "real patients (of 18) for whom the SDMetrics attacker could make any prediction, per attack, sensitive "
     "field and run. Patients without a prediction were scored as protected."),
    ("S_privacy_revised_attack_by_target", "Supplementary Figure S6. Revised attribute inference attack: "
     "percentage of the 18 real patients whose target value was correctly inferred by the strongest attacker "
     "in each run (mean and range across 10 runs), compared with a prior-only guess that ignores the "
     "attacker's knowledge (black bar)."),
    ("S_privacy_revised_lift_by_seed", "Supplementary Figure S7. Attack success above the prior-only guess, "
     "averaged over eight targets, for each of 10 runs per synthesizer. Values ≤ 0 indicate that the synthetic "
     "data gave the attacker no information beyond guessing."),
    ("S_holdout_membership", "Supplementary Figure S8. Memorization and membership inference with synthesizers "
     "retrained in 6-fold cross-validation (2 repetitions). Left: similarity of each real patient to the "
     "closest synthetic record when the patient was in the synthesizer's training data (colored) or held out "
     "(gray). Right: area under the ROC curve for separating training from held-out patients."),
    ("S_holdout_tstr", "Supplementary Figure S9. Train-on-synthetic, test-on-real (TSTR) and train-on-real, "
     "test-on-real (TRTR) discrimination for two endpoints (postoperative complication, 6/18 events; "
     "radiotherapy, 12/18 events), with bootstrap 95% CIs. Dashed line: chance."),
]


# ---------- document helpers ----------

def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_runs(paragraph, text):
    """Add text with **bold** spans."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def add_markdown(doc, path):
    for line in open(path, encoding="utf-8").read().splitlines():
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_runs(doc.add_paragraph(), line)


def add_table(doc, title, header, rows, widths, note=None, font=8):
    heading = doc.add_paragraph()
    heading.add_run(title).bold = True
    heading.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text, w in zip(table.rows[0].cells, header, widths):
        cell.width = Inches(w)
        shade(cell, "E4E3DF")
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(font)
    for row in rows:
        cells = table.add_row().cells
        for cell, text, w in zip(cells, row, widths):
            cell.width = Inches(w)
            cell.paragraphs[0].add_run(str(text)).font.size = Pt(font)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    if note:
        r = p.add_run(note)
        r.italic = True
        r.font.size = Pt(8)


def fmt_p(p):
    return "< 0.001" if p < 0.001 else f"{p:.3f}"


# ---------- tables from analysis outputs ----------

def table_missingness(doc):
    t = pd.read_csv(OUT / "missingness_table.csv")
    t = t[t.Missing > 0]
    rows = [(r.Variable, r.Type, f"{r.Missing}/18 ({r.Missing_pct:.1f}%)", r.Imputation) for r in t.itertuples()]
    add_table(doc, "Supplementary Table S1. Variable-level missingness before imputation",
              ["Variable", "Type", "Missing", "Imputation"], rows, [3.0, 1.0, 1.2, 1.6],
              note="Variables not listed had no missing values (21 of 61). Blank cells and cells recorded as "
                   "'Missing' in the source database were counted as missing.")


def table_safe_harbor(doc):
    add_table(doc, "Supplementary Table S2. Audit of study variables against the 18 HIPAA Safe Harbor identifier "
                   "categories (45 CFR §164.514(b)(2))",
              ["Identifier category", "Present in study data", "Handling"], SAFE_HARBOR, [2.0, 2.1, 2.7],
              note="This audit documents how each identifier category was handled; it is not a formal "
                   "Safe Harbor or Expert Determination certification of the released synthetic datasets.")


def table_fidelity(doc):
    s = pd.read_csv(OUT / "fidelity_summary.csv").set_index(["Source", "Metric"])
    rows = []
    for src in SOURCES:
        row = [src]
        for m in ["Overall", "Column Shapes", "Column Pair Trends"]:
            r = s.loc[(src, m)]
            row.append(f"{r.Mean:.1f} ± {r.SD:.1f} ({r.Min:.1f}–{r.Max:.1f})")
        rows.append(row)
    tests = pd.read_csv(OUT / "fidelity_tests.csv")
    kw = tests[tests.Comparison.str.startswith("Kruskal")].set_index("Metric").p
    add_table(doc, "Supplementary Table S3. SDMetrics fidelity across 10 seeded runs per synthesizer",
              ["Source", "Overall, %", "Column shapes, %", "Column pair trends, %"], rows, [1.5, 1.8, 1.8, 1.8],
              note="Mean ± SD (range) across 10 runs. Kruskal–Wallis across the four synthesizers: overall "
                   f"p {fmt_p(kw['Overall'])}, column shapes p {fmt_p(kw['Column Shapes'])}, column pair trends "
                   f"p {fmt_p(kw['Column Pair Trends'])}. Holm-adjusted pairwise comparisons are provided in "
                   "the analysis repository (fidelity_tests.csv).")


def table_associations(doc):
    a = pd.read_csv(OUT / "utility_associations_summary.csv")
    order = list(dict.fromkeys(a.Association))
    rows = []
    for assoc in order:
        g = a[a.Association == assoc].set_index("Source")
        r0 = g.iloc[0]
        row = [assoc, f"{r0.Real_logOR:.2f} ({r0.Real_CI_low:.2f} to {r0.Real_CI_high:.2f})"]
        for src in SOURCES:
            r = g.loc[src]
            row.append(f"{r.logOR_mean:.2f} [{r.Same_direction_runs_pct:.0f}%]")
        rows.append(row)
    add_table(doc, "Supplementary Table S4. Reproduction of prespecified clinical associations",
              ["Association", "Real log OR (95% CI)"] + [SHORT_SOURCE.get(s, s) for s in SOURCES], rows,
              [1.75, 1.15] + [0.8] * 5, font=7,
              note="Synthetic entries: mean log odds ratio across 10 runs [percentage of runs with the same "
                   "direction as the real estimate]. Log odds ratios are Haldane–Anscombe corrected.")


def table_global(doc):
    g = pd.read_csv(OUT / "utility_global_summary.csv").set_index(["Source", "Metric"])
    rows = []
    for src in SOURCES:
        c = g.loc[(src, "Corr_of_correlations")]
        d = g.loc[(src, "Mean_abs_diff")]
        s = g.loc[(src, "Sign_agreement_strong_pct")]
        rows.append([src, f"{c.Mean:.2f} ({c.CI95_low:.2f} to {c.CI95_high:.2f})",
                     f"{d.Mean:.3f} ({d.CI95_low:.3f} to {d.CI95_high:.3f})",
                     f"{s.Mean:.1f} ({s.CI95_low:.1f} to {s.CI95_high:.1f})"])
    add_table(doc, "Supplementary Table S5. Preservation of the global correlation structure (55 variables, "
                   "1,485 pairs)",
              ["Source", "Correlation of correlations", "Mean absolute difference", "Sign agreement, strong pairs, %"],
              rows, [1.5, 1.8, 1.8, 1.8], note="Mean (95% CI) across 10 runs. Strong pairs: |ρ| ≥ 0.3 in the "
                                               "real data (410 pairs).")


def table_original_attack(doc):
    s = pd.read_csv(OUT / "privacy_original_summary.csv").set_index("Source")
    r = pd.read_csv(OUT / "privacy_original_risk_by_seed.csv").set_index("Source")
    rows = [[src, f"{s.loc[src].Evaluable_mean:.1f} (max {s.loc[src].Evaluable_max:.0f})",
             f"{r.loc[src]['mean']:.1f} ({r.loc[src]['min']:.1f}–{r.loc[src]['max']:.1f})"] for src in SOURCES]
    add_table(doc, "Supplementary Table S6. Re-analysis of the original SDMetrics attack specification across "
                   "10 runs",
              ["Source", "Patients evaluable per attack, of 18", "Reported risk per run, % mean (range)"], rows,
              [1.8, 2.4, 2.6],
              note="Original key fields and clinical sensitive fields; kNN, naive Bayes and random-forest "
                   "SDMetrics attackers. Reported risk = 1 − SDMetrics privacy score, averaged over sensitive "
                   "fields and attackers within each run.")


def table_revised_by_target(doc):
    t = pd.read_csv(OUT / "privacy_revised_by_target.csv")
    targets = list(dict.fromkeys(t.Target))
    rows = []
    for tgt in targets:
        g = t[t.Target == tgt].set_index("Source")
        row = [tgt, f"{g.iloc[0].Prior_correct_of_18:g}"]
        for src in SOURCES:
            r = g.loc[src]
            row.append(f"{r.Median_correct_of_18}/18 ({round(r.Mean_lift_pp, 1) + 0.0:+.1f})")
        rows.append(row)
    add_table(doc, "Supplementary Table S7. Revised attribute inference attack by target (strongest attacker per run)",
              ["Target", "Prior-only guess, correct of 18"] + [SHORT_SOURCE.get(s, s) for s in SOURCES], rows,
              [1.55, 0.9] + [0.87] * 5, font=7,
              note="Median number of real patients correctly inferred across 10 runs (mean lift over the "
                   "prior-only guess, percentage points). Continuous targets: correct if within ±10% of the "
                   "real interquartile range. Exact 95% CIs for a count k of 18 are wide (for example, 15/18: "
                   "58.6–96.4%); per-run values and CIs are in privacy_revised_by_target.csv.")


def table_revised_by_model(doc):
    m = pd.read_csv(OUT / "privacy_revised_by_model.csv").set_index(["Source", "Metric"])
    tests = pd.read_csv(OUT / "privacy_revised_tests.csv")
    vs_base = {c.split(" vs ")[0]: p for c, p in zip(tests.Comparison, tests.p_Holm) if c.endswith(BASELINE)}
    rows = []
    for src in SOURCES:
        s = m.loc[(src, "Mean_success_pct")]
        lift = m.loc[(src, "Mean_lift_pp")]
        rows.append([src, f"{s.Mean:.1f} ({s.CI95_low:.1f} to {s.CI95_high:.1f})",
                     f"{lift.Mean:+.1f} ({lift.CI95_low:+.1f} to {lift.CI95_high:+.1f})",
                     "—" if src == BASELINE else fmt_p(vs_base[src])])
    kw = tests[tests.Comparison.str.startswith("Kruskal")].p.iloc[0]
    add_table(doc, "Supplementary Table S8. Revised attribute inference attack by synthesizer",
              ["Source", "Mean success, % (95% CI)", "Lift over prior-only guess, pp (95% CI)",
               "Holm p vs baseline"], rows, [1.5, 1.8, 2.1, 1.3],
              note=f"Averaged over eight targets per run (strongest attacker); 95% CI across 10 runs. "
                   f"Kruskal–Wallis across synthesizers p {fmt_p(kw)}.")


def table_holdout(doc):
    mem = pd.read_csv(OUT / "holdout_membership_summary.csv").set_index("Source")
    aia = pd.read_csv(OUT / "holdout_aia_member_vs_nonmember.csv").set_index("Source")
    rows = []
    for src in SOURCES:
        a, b = mem.loc[src], aia.loc[src]
        rows.append([src, f"{a['mean']:.2f} ({a['min']:.2f}–{a['max']:.2f})",
                     f"{100 * a.Median_similarity_members:.1f} / {100 * a.Median_similarity_nonmembers:.1f}",
                     f"{b.Success_members_pct:.1f} / {b.Success_nonmembers_pct:.1f}",
                     f"{b.Member_minus_nonmember_pp:+.1f} ({b.CI95_low:+.1f} to {b.CI95_high:+.1f})"])
    add_table(doc, "Supplementary Table S9. Memorization and membership analyses (6-fold cross-validation, "
                   "2 repetitions)",
              ["Source", "Membership AUC, mean (range)", "Closest-record similarity, %, members / held-out",
               "Attack success, %, members / held-out", "Difference in lift, pp (95% CI)"], rows,
              [1.4, 1.3, 1.5, 1.4, 1.4], font=7,
              note="Members: patients in the synthesizer's training folds; held-out: patients excluded from "
                   "training. The independent-marginals baseline resamples exact values from training patients, "
                   "so its AUC is the reference level for this measure.")


def table_tstr(doc):
    t = pd.read_csv(OUT / "holdout_tstr_summary.csv").set_index(["Source", "Endpoint"])
    rows = []
    for src in ["Real (TRTR)"] + SOURCES:
        rows.append([src] + [f"{t.loc[(src, e)].AUC:.2f} ({t.loc[(src, e)].CI_low:.2f}–{t.loc[(src, e)].CI_high:.2f})"
                             for e in ["Postop complication", "Radiotherapy (Y/N)"]])
    add_table(doc, "Supplementary Table S10. Train-on-synthetic, test-on-real discrimination",
              ["Training data", "Postoperative complication (6/18), AUC (95% CI)",
               "Radiotherapy (12/18), AUC (95% CI)"], rows, [1.8, 2.5, 2.5],
              note="Random forest, 34 pre-/intraoperative predictors; predictions for held-out patients pooled "
                   "over 6 folds and averaged over 2 repetitions; bootstrap 95% CI.")


def main():
    doc = Document()
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        for side in ["left_margin", "right_margin", "top_margin", "bottom_margin"]:
            setattr(section, side, Inches(0.8))
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for level in (1, 2, 3):
        h = doc.styles[f"Heading {level}"]
        h.font.name = "Arial"
        # Built-in heading styles use theme fonts, which override font.name.
        fonts = h.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            fonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(attr), "Arial")
        h.font.color.rgb = RGBColor(0x0B, 0x0B, 0x0B)
        h.font.size = Pt({1: 14, 2: 12, 3: 10.5}[level])

    title = doc.add_paragraph()
    title.add_run("Supplementary Material").bold = True
    title.runs[0].font.size = Pt(16)
    doc.add_paragraph("Cybersecurity and AI in Neuro-oncology: An Exploratory Evaluation of Synthetic Data "
                      "Privacy and Fidelity in a Rare Brain Tumor Cohort")

    add_markdown(doc, OUT / "Supplementary_Methods_draft.md")
    doc.add_page_break()
    add_markdown(doc, OUT / "Supplementary_Results_draft.md")

    doc.add_heading("Supplementary Tables", level=1).paragraph_format.page_break_before = True
    table_missingness(doc)
    table_safe_harbor(doc)
    table_fidelity(doc)
    table_associations(doc)
    table_global(doc)
    table_original_attack(doc)
    table_revised_by_target(doc)
    table_revised_by_model(doc)
    table_holdout(doc)
    table_tstr(doc)

    doc.add_heading("Supplementary Figures", level=1).paragraph_format.page_break_before = True
    for name, caption in FIGURES:
        path = FIG_DIR / f"{name}.png"
        with Image.open(path) as im:
            aspect = im.height / im.width
        # Cap height so each figure fits on a page with its caption.
        doc.add_picture(str(path), width=Inches(min(6.6, MAX_FIG_HEIGHT / aspect)))
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(18)
        add_runs(p, caption)
        p.runs[0].font.size = Pt(9)

    out = OUT / "Supplementary_Material_revision.docx"
    doc.save(out)
    print(f"saved {out}")


if __name__ == "__main__":
    main()
